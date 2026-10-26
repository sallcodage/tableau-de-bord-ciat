"""Exports des analyses : Excel (openpyxl) et PDF (reportlab + graphiques matplotlib)."""
import io
from datetime import date

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.chart import BarChart, Reference  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from reportlab.lib import colors  # noqa: E402
from reportlab.lib.pagesizes import A4  # noqa: E402
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # noqa: E402
from reportlab.lib.units import cm  # noqa: E402
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

from .analyses import faits_marquants, fcfa, fr, indicateurs, projets_en_retard, series  # noqa: E402
from .models import DOMAINES_COURTS, Jalon  # noqa: E402

BRUN, CARAMEL, CREME = "#4B2E1F", "#C8956D", "#F3E6DA"

LIGNES_KPI = [
    ("Idées reçues", "idees", ""), ("Projets présélectionnés", "preselectionnes", ""),
    ("Projets incubés", "selectionnes", ""), ("Prototypes réalisés", "prototypes", ""),
    ("Projets sortis (lancés)", "sortis", ""), ("Abandons", "abandons", ""),
    ("Taux de sélection", "taux_selection", "%"), ("Taux de prototypes (parmi les incubés)", "taux_prototype", "%"),
    ("Taux de réussite en incubation", "taux_reussite", "%"), ("Taux de jalons en retard", "taux_retard", "%"),
    ("Projets avec un jalon en retard", "projets_en_retard", ""),
    ("Financements obtenus (FCFA)", "montant_obtenu", "fcfa"), ("Projets financés", "projets_finances", ""),
    ("Partenariats actifs", "partenariats_actifs", ""), ("Part des femmes porteuses", "part_femmes", "%"),
    ("Emplois créés", "emplois_crees", ""),
]


def _val(k, cle, unite):
    v = k[cle]
    if unite == "%":
        return f"{fr(v)} %"
    if unite == "fcfa":
        return fcfa(v)
    return fr(v)


# --------------------------------------------------------------------------- Excel
def excel(filtres, voit_noms):
    projets = filtres.projets()
    k, s = indicateurs(projets), series(projets)
    wb = Workbook()
    entete = Font(bold=True, color="FFFFFF")
    fond = PatternFill("solid", fgColor=BRUN[1:])

    def feuille(titre, colonnes, lignes, ws=None):
        ws = ws or wb.create_sheet(titre)
        ws.append(colonnes)
        for c in ws[ws.max_row]:
            c.font, c.fill = entete, fond
        for ligne in lignes:
            ws.append(ligne)
        for i, col in enumerate(colonnes, start=1):
            largeur = max([len(str(col))] + [len(str(l[i - 1])) for l in lignes] or [10])
            ws.column_dimensions[get_column_letter(i)].width = min(45, largeur + 2)
        ws.freeze_panes = ws.cell(row=ws.max_row - len(lignes) + 1, column=1)
        return ws

    ws = wb.active
    ws.title = "Synthèse"
    ws.append(["Tableau de bord décisionnel du CIAT — ISEP Amadou Trawaré"])
    ws["A1"].font = Font(bold=True, size=14, color=BRUN[1:])
    ws.append([f"Périmètre : {filtres.description()}"])
    ws.append([f"Export du {date.today():%d/%m/%Y} — données fictives de démonstration"])
    ws.append([])
    feuille("Synthèse", ["Indicateur", "Valeur"], [[lib, k[cle]] for lib, cle, _ in LIGNES_KPI], ws=ws)
    ws.append([])
    ws.append(["Faits marquants"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    for fait in faits_marquants(projets):
        ws.append([fait])
    ws.column_dimensions["A"].width = 60

    wc = feuille("Cohortes", ["Cohorte", "Idées", "Incubés", "Prototypes", "Sortis", "Abandons", "Taux sélection %",
                              "Taux prototypes %", "Taux réussite %", "Part femmes %", "Financements FCFA", "Taux retard %"],
                 [[c["cohorte"], c["idees"], c["selectionnes"], c["prototypes"], c["sortis"], c["abandons"],
                   c["taux_selection"] if c["taux_selection"] is not None else "",
                   c["taux_prototype"] if c["taux_prototype"] is not None else "",
                   c["taux_reussite"] if c["taux_reussite"] is not None else "en cours",
                   c["part_femmes"], c["montant_obtenu"], c["taux_retard"]] for c in s["cohortes"]])
    if s["cohortes"]:  # graphique Excel natif
        g = BarChart()
        g.title, g.y_axis.title = "Comparaison des cohortes", "Projets"
        n = len(s["cohortes"]) + 1
        g.add_data(Reference(wc, min_col=2, max_col=5, min_row=1, max_row=n), titles_from_data=True)
        g.set_categories(Reference(wc, min_col=1, min_row=2, max_row=n))
        g.width, g.height = 18, 9
        wc.add_chart(g, f"A{n + 3}")

    feuille("Domaines", ["Domaine", "Idées", "Incubés", "Taux de sélection %"],
            [[d["domaine"], d["idees"], d["selectionnes"], d["taux_selection"]] for d in s["domaines"]])
    feuille("Filières ISEP", ["Filière", "Département", "Idées", "Incubés", "Sortis"],
            [[f["filiere"], f["departement"], f["idees"], f["selectionnes"], f["sortis"]] for f in s["filieres"]])
    feuille("Genre", ["Genre", "Idées", "Incubés", "Sortis", "Taux de sélection %"],
            [[g["genre"], g["idees"], g["selectionnes"], g["sortis"], g["taux_selection"]] for g in s["genre"]])

    colonnes = ["Code", "Titre", "Domaine", "Cohorte", "Porteur", "Genre", "Région", "Filière ISEP", "Dépôt", "Présélectionné",
                "Incubé", "Note comité", "Prototype", "Statut", "Progression %", "Emplois"]
    feuille("Projets", colonnes, [[
        p.code, p.titre, p.domaine_court, p.cohorte.annee,
        f"{p.porteur.prenom} {p.porteur.nom}" if voit_noms else p.porteur.code,
        p.porteur.get_genre_display(), p.porteur.region, p.porteur.get_filiere_display() or "Externe", p.date_depot, "oui" if p.preselectionne else "non",
        "oui" if p.selectionne else "non", p.score_comite or "", p.date_prototype or "", p.get_statut_display(),
        p.progression, p.emplois_crees] for p in projets.order_by("code")])

    jalons = Jalon.objects.filter(projet__in=projets).select_related("projet")
    feuille("Jalons", ["Projet", "Jalon", "Échéance", "Réalisé le", "État", "Jours de retard"],
            [[j.projet.code, j.titre, j.echeance, j.date_realisation or "", Jalon.Etat(j.etat()).label,
              j.jours_retard if j.etat() in ("en_retard", "realise_retard") else 0]
             for j in jalons.order_by("projet__code", "echeance")])

    from .models import Financement
    feuille("Financements", ["Projet", "Organisme", "Type", "Montant FCFA", "Date", "Statut"],
            [[f.projet.code, f.organisme, f.get_type_display(), f.montant, f.date, f.get_statut_display()]
             for f in Financement.objects.filter(projet__in=projets).select_related("projet")])
    for w in wb.worksheets:
        for ligne in w.iter_rows():
            for c in ligne:
                c.alignment = Alignment(vertical="top")
    sortie = io.BytesIO()
    wb.save(sortie)
    return sortie.getvalue()


# --------------------------------------------------------------------------- PDF
def _graphique(fig):
    tampon = io.BytesIO()
    fig.savefig(tampon, format="png", dpi=160, bbox_inches="tight")
    plt.close(fig)
    tampon.seek(0)
    return tampon


def _style_axes(ax):
    for cote in ("top", "right"):
        ax.spines[cote].set_visible(False)
    ax.spines["left"].set_color("#BBB")
    ax.spines["bottom"].set_color("#BBB")
    ax.tick_params(colors="#555", labelsize=8)


def graphique_entonnoir(s):
    fig, ax = plt.subplots(figsize=(6.2, 2.4))
    etapes = [e["etape"] for e in s["entonnoir"]][::-1]
    n = [e["n"] for e in s["entonnoir"]][::-1]
    ax.barh(etapes, n, color=[BRUN] + [CARAMEL] * 4)  # liste inversée : la barre foncée = projets sortis
    for i, v in enumerate(n):
        ax.text(v + max(n) * 0.01, i, str(v), va="center", fontsize=8, color="#333")
    ax.set_title("Parcours des projets (entonnoir)", fontsize=10, loc="left", color=BRUN)
    _style_axes(ax)
    return _graphique(fig)


def graphique_cohortes(s):
    fig, ax = plt.subplots(figsize=(6.2, 2.6))
    cs = [str(c["cohorte"]) for c in s["cohortes"]]
    largeur = 0.2
    for i, (cle, lib, coul) in enumerate([("idees", "Idées", "#E3CDB8"), ("selectionnes", "Incubés", CARAMEL),
                                          ("prototypes", "Prototypes", "#8C5A3C"), ("sortis", "Sortis", BRUN)]):
        ax.bar([x + (i - 1.5) * largeur for x in range(len(cs))], [c[cle] for c in s["cohortes"]], largeur,
               label=lib, color=coul)
    ax.set_xticks(range(len(cs)), cs)
    ax.legend(fontsize=7, frameon=False, ncol=4, loc="lower right", bbox_to_anchor=(1, 1.0))
    ax.set_title("Comparaison des cohortes", fontsize=10, loc="left", color=BRUN)
    _style_axes(ax)
    return _graphique(fig)


def graphique_domaines(s):
    fig, ax = plt.subplots(figsize=(6.2, 2.4))
    ds = s["domaines"][::-1]
    ax.barh([d["domaine"] for d in ds], [d["idees"] for d in ds], color="#E3CDB8", label="Idées")
    ax.barh([d["domaine"] for d in ds], [d["selectionnes"] for d in ds], color=BRUN, label="Incubés")
    ax.legend(fontsize=7, frameon=False)
    ax.set_title("Idées et projets incubés par domaine", fontsize=10, loc="left", color=BRUN)
    _style_axes(ax)
    return _graphique(fig)


def pdf(filtres, titre="Fiche de synthèse", synthese_ia=None):
    projets = filtres.projets()
    k, s = indicateurs(projets), series(projets)
    tampon = io.BytesIO()
    doc = SimpleDocTemplate(tampon, pagesize=A4, leftMargin=1.8 * cm, rightMargin=1.8 * cm,
                            topMargin=1.6 * cm, bottomMargin=1.6 * cm, title=titre, author="CIAT — ISEP Amadou Trawaré")
    st = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=st["Heading1"], textColor=colors.HexColor(BRUN), fontSize=17, spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=st["Heading2"], textColor=colors.HexColor(BRUN), fontSize=12, spaceBefore=10)
    txt = ParagraphStyle("t", parent=st["BodyText"], fontSize=9, leading=12.5)
    petit = ParagraphStyle("p", parent=txt, fontSize=7.5, textColor=colors.HexColor("#6F6158"))

    elements = [
        Paragraph("CIAT — ISEP Amadou Trawaré de Diamniadio", petit),
        Paragraph(titre, h1),
        Paragraph(f"Périmètre : <b>{filtres.description()}</b> · édité le {date.today():%d/%m/%Y}", txt),
        Spacer(1, 8),
    ]
    tuiles = [[Paragraph(f"<font size=15 color='{BRUN}'><b>{v}</b></font><br/><font size=7.5>{lib}</font>", txt)
               for lib, v in [("Idées reçues", k["idees"]), ("Projets incubés", k["selectionnes"]),
                              ("Prototypes", k["prototypes"]), ("Taux de réussite", f"{fr(k['taux_reussite'])} %")]],
              [Paragraph(f"<font size=15 color='{BRUN}'><b>{v}</b></font><br/><font size=7.5>{lib}</font>", txt)
               for lib, v in [("Financements obtenus", fcfa(k["montant_obtenu"])), ("Jalons en retard", f"{fr(k['taux_retard'])} %"),
                              ("Part des femmes", f"{fr(k['part_femmes'])} %"), ("Emplois créés", k["emplois_crees"])]]]
    t = Table(tuiles, colWidths=[4.3 * cm] * 4)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(CREME)),
                           ("BOX", (0, 0), (-1, -1), 0.5, colors.white), ("INNERGRID", (0, 0), (-1, -1), 3, colors.white),
                           ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]))
    elements += [t, Paragraph("Faits marquants", h2)]
    elements += [Paragraph(f"• {f}", txt) for f in faits_marquants(projets)]
    if synthese_ia:
        elements += [Paragraph("Note de synthèse (rédigée par IA à partir des chiffres ci-dessus, à valider)", h2)]
        elements += [Paragraph(p.strip(), txt) for p in synthese_ia.split("\n") if p.strip()]
    if k["idees"]:
        elements += [Spacer(1, 6), Image(graphique_entonnoir(s), width=16 * cm, height=6.2 * cm)]
        if len(s["cohortes"]) > 1:
            elements += [Image(graphique_cohortes(s), width=16 * cm, height=6.7 * cm)]
        elements += [Image(graphique_domaines(s), width=16 * cm, height=6.2 * cm)]

    elements.append(Paragraph("Indicateurs détaillés", h2))
    lignes = [["Indicateur", "Valeur"]] + [[lib, _val(k, cle, u)] for lib, cle, u in LIGNES_KPI]
    t = Table(lignes, colWidths=[11 * cm, 5 * cm])
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(BRUN)),
                           ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                           ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FAF7F3")]),
                           ("ALIGN", (1, 0), (1, -1), "RIGHT"), ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E8DFD6"))]))
    elements.append(t)

    retards = projets_en_retard(projets)
    if retards:
        elements.append(Paragraph("Projets en incubation avec un jalon en retard", h2))
        lignes = [["Projet", "Domaine", "Jalon le plus en retard", "Retard"]] + [
            [r["projet"].code, r["projet"].domaine_court, r["jalon"], f"{r['jours']} j"] for r in retards]
        t = Table(lignes, colWidths=[3.6 * cm, 3.2 * cm, 6.6 * cm, 2.6 * cm])
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(BRUN)),
                               ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                               ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E8DFD6"))]))
        elements.append(t)
    elements += [Spacer(1, 10), Paragraph("Données fictives de démonstration — évaluation pratique IA générative, ISEP Amadou Trawaré.", petit)]
    doc.build(elements)
    return tampon.getvalue()
