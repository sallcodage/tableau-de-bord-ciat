"""
Tests des fonctions essentielles. Lancer : python manage.py test
Les appels à Gemini sont simulés : aucun quota consommé.
"""
import io
from datetime import date
from unittest.mock import patch

from django.core.management import call_command
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import User

from . import assistant
from .analyses import Filtres, indicateurs, series
from .models import Financement, Jalon, Projet
from .qualite import controles


class Donnees(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("seed_demo", stdout=io.StringIO())
        cls.direction = User.objects.get(username="demo_direction")
        cls.coord = User.objects.get(username="demo_coordonnateur")


class IndicateursTests(Donnees):
    def test_entonnoir_coherent(self):
        k = indicateurs(Projet.objects.all())
        self.assertGreaterEqual(k["idees"], k["preselectionnes"])
        self.assertGreaterEqual(k["selectionnes"], k["prototypes"])
        self.assertGreaterEqual(k["prototypes"], 0)
        self.assertEqual(k["taux_selection"], round(100 * k["selectionnes"] / k["idees"], 1))

    def test_filtre_cohorte_et_domaine(self):
        f = Filtres(cohortes=[2025], domaines=["agriculture"])
        self.assertTrue(all(p.cohorte.annee == 2025 and p.domaine == "agriculture" for p in f.projets()))

    def test_filtre_filiere_et_departement(self):
        apd = Filtres(filiere="apd").projets()
        self.assertTrue(apd.exists())
        self.assertTrue(all(p.porteur.filiere == "apd" for p in apd))
        auto = Filtres(filiere="dep:Automobile").projets()
        self.assertTrue(all(p.porteur.departement == "Département Automobile" for p in auto))
        # les nouvelles filières n'existent que pour l'appel 2027
        self.assertFalse(Projet.objects.filter(porteur__filiere__in=["ia", "cyber", "efficacite", "biomed"])
                         .exclude(cohorte__annee=2027).exists())

    def test_taux_non_significatifs_a_none(self):
        c = {x["cohorte"]: x for x in series(Projet.objects.all())["cohortes"]}
        self.assertIsNone(c[2027]["taux_selection"])     # appel en cours d'évaluation
        self.assertIsNone(c[2026]["taux_reussite"])      # incubation en cours
        self.assertIsNotNone(c[2025]["taux_reussite"])

    def test_etat_des_jalons(self):
        p = Projet.objects.filter(statut="incubation").first()
        j = Jalon(projet=p, titre="t", echeance=date(2026, 1, 1))
        self.assertEqual(j.etat(date(2026, 2, 1)), Jalon.Etat.EN_RETARD)
        self.assertEqual(j.etat(date(2025, 12, 1)), Jalon.Etat.A_VENIR)
        j.date_realisation = date(2026, 1, 10)
        self.assertEqual(j.etat(), Jalon.Etat.REALISE_RETARD)

    def test_controle_qualite_detecte_les_anomalies_volontaires(self):
        avec_anomalie = {c["regle"] for c in controles() if c["nb"]}
        self.assertIn("Projet sélectionné sans note du comité", avec_anomalie)
        self.assertIn("Financement daté avant le dépôt du projet", avec_anomalie)
        self.assertIn("Prototype déclaré sans sélection en incubation", avec_anomalie)


@override_settings(GEMINI_API_KEY="")
class AssistantReglesTests(Donnees):
    def test_question_groupee(self):
        r = assistant.repondre("Combien de projets incubés par domaine en 2025 ?")
        self.assertTrue(r["succes"])
        total = sum(l["valeur"] for l in r["resultat"]["lignes"])
        self.assertEqual(total, Projet.objects.filter(cohorte__annee=2025, selectionne=True).count())

    def test_question_hors_perimetre(self):
        self.assertFalse(assistant.repondre("Quel temps fait-il à Dakar ?")["succes"])


@override_settings(GEMINI_API_KEY="cle-test")
class AssistantGeminiTests(Donnees):
    @patch("pilotage.assistant.analyser_gemini")
    def test_specification_validee_par_liste_blanche(self, mock):
        mock.return_value = {"mesure": "nb_idees; DROP TABLE", "grouper_par": None, "filtres": {}}
        r = assistant.repondre("question piégée")
        self.assertFalse(r["succes"])

    @patch("pilotage.assistant.analyser_gemini")
    def test_filtres_inconnus_ignores(self, mock):
        mock.return_value = {"mesure": "nb_idees", "grouper_par": "cohorte",
                             "filtres": {"domaines": ["inexistant"], "cohortes": [1999]}}
        r = assistant.repondre("idées par cohorte")
        self.assertTrue(r["succes"])
        self.assertEqual(r["spec"]["filtres"]["domaines"], [])
        self.assertEqual(r["spec"]["filtres"]["cohortes"], [])

    @patch("pilotage.assistant.analyser_gemini", side_effect=RuntimeError("quota"))
    def test_panne_ia_repli_sur_regles(self, _):
        r = assistant.repondre("Nombre de prototypes en agriculture")
        self.assertTrue(r["succes"])
        self.assertEqual(r["mode"], "regles")


class VuesEtDroitsTests(Donnees):
    def test_connexion_obligatoire(self):
        self.assertEqual(self.client.get(reverse("pilotage:dashboard")).status_code, 302)

    def test_pages_principales(self):
        self.client.force_login(self.coord)
        for url in [reverse("pilotage:dashboard"), reverse("pilotage:dashboard") + "?cohorte=2025&genre=F",
                    reverse("pilotage:projets"), reverse("pilotage:cohortes"),
                    reverse("pilotage:fiche_cohorte", args=[2025]), reverse("pilotage:assistant"),
                    reverse("pilotage:donnees"), reverse("pilotage:methodologie"),
                    reverse("pilotage:fiche_projet", args=[Projet.objects.first().code])]:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 200)

    def test_direction_ne_voit_pas_les_noms_ni_l_import(self):
        self.client.force_login(self.direction)
        p = Projet.objects.select_related("porteur").first()
        page = self.client.get(reverse("pilotage:projets") + f"?cohorte={p.cohorte.annee}")
        self.assertNotContains(page, p.porteur.nom + "<")
        self.assertEqual(self.client.get(reverse("pilotage:donnees")).status_code, 403)
        # le filtre par nom est ignoré pour la direction
        f = Filtres.depuis_requete(self.client.get("/").wsgi_request.GET.__class__("porteur=Diop"), voit_noms=False)
        self.assertEqual(f.porteur, "")

    def test_export_excel_anonymise_pour_la_direction(self):
        self.client.force_login(self.direction)
        r = self.client.get(reverse("pilotage:export_excel") + "?cohorte=2025")
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(io.BytesIO(r.content))
        self.assertIn("Projets", wb.sheetnames)
        porteurs = [c.value for c in wb["Projets"]["E"][1:]]
        self.assertTrue(all(str(v).startswith("P-") for v in porteurs))

    def test_export_pdf(self):
        self.client.force_login(self.direction)
        r = self.client.get(reverse("pilotage:export_pdf") + "?domaine=tic")
        self.assertEqual(r["Content-Type"], "application/pdf")
        self.assertTrue(r.content.startswith(b"%PDF"))

    def test_import_csv_invalide_annule_tout(self):
        self.client.force_login(self.coord)
        avant = Financement.objects.count()
        mauvais = io.BytesIO("projet;organisme;type;montant;date;statut\nINCONNU;X;prix;100;2025-01-01;obtenu\n".encode())
        mauvais.name = "financements.csv"
        self.client.post(reverse("pilotage:donnees"), {"financements": mauvais})
        self.assertEqual(Financement.objects.count(), avant)

    @override_settings(GEMINI_API_KEY="")
    def test_journal_des_requetes(self):
        self.client.force_login(self.direction)
        self.client.post(reverse("pilotage:assistant"), {"question": "Taux de réussite par cohorte"})
        self.assertEqual(self.direction.requetes.count(), 1)
